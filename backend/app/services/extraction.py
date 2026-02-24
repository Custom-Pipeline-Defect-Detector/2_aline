from pathlib import Path
from typing import Tuple
import pdfplumber
from docx import Document as DocxDocument
import openpyxl


def extract_text(file_path: Path) -> Tuple[str, str]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        text = []
        with pdfplumber.open(file_path) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                page_text = page.extract_text() or ""
                text.append(f"[page:{index}]\n{page_text}")
        return "\n".join(text), "pdf"
    if suffix in {".docx", ".doc"}:
        doc = DocxDocument(file_path)
        text = [paragraph.text for paragraph in doc.paragraphs]
        return "\n".join(text), "docx"
    if suffix in {".xlsx", ".xls"}:
        book = openpyxl.load_workbook(file_path, data_only=True)
        lines = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = ", ".join([str(cell) for cell in row if cell is not None])
                if row_text:
                    lines.append(f"[sheet:{sheet.title}] {row_text}")
        return "\n".join(lines), "xlsx"
    try:
        return file_path.read_text(encoding="utf-8"), "text"
    except UnicodeDecodeError:
        return file_path.read_text(encoding="latin-1"), "text"
